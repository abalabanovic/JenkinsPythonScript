import pandas as pd
import re
import openpyxl
from openpyxl.utils import get_column_letter

# Define input and output file names
#input_file =   Replace with the actual input file
output_file = "output.xlsx"

# Initialize lists for extracted data
urls = []
build_dates = []

# Read and process the input file
with open(input_file, "r", encoding="utf-8") as file:
    lines = file.readlines()
    
    for i in range(len(lines) - 1):  # Loop through lines
        url = lines[i].strip()
        build_date_match = re.search(r"Last Build:\s*([\d-]+\s[\d:]+)", lines[i + 1].strip())


        if build_date_match:
            urls.append(url)  # Store URL
            build_dates.append(build_date_match.group(1).strip())  # Store build date

# Create a DataFrame
df = pd.DataFrame({'Pipeline URL': urls, 'Build Date': build_dates})

# Save to an Excel file
df.to_excel(output_file, index=False)

wb = openpyxl.load_workbook(output_file)
ws = wb.active

# Adjust column widths based on content
for col_num in range(1, len(df.columns) + 1):  # Loop through columns
    max_length = 0
    column = get_column_letter(col_num)
    
    # Loop through all rows in the column to find the longest value
    for row in ws.iter_rows(min_col=col_num, max_col=col_num):
        for cell in row:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
    
    # Set column width with a little extra buffer
    adjusted_width = (max_length + 2)
    ws.column_dimensions[column].width = adjusted_width

# Save the updated Excel file with adjusted column widths
wb.save(output_file)

print(f"✅ Successfully saved to {output_file}")
